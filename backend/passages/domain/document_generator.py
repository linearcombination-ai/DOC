import json
import time
from typing import TYPE_CHECKING, Mapping, Optional, Sequence

from celery import current_task
from doc.config import settings
from doc.domain import worker
from doc.domain.bible_books import BOOK_NAMES
from doc.domain.email_utils import send_email_with_attachment, should_send_email
from doc.domain.model import Attachment, USFMBook
from doc.domain.parsing import (
    usfm_book_content,
    split_into_verses,
)
from doc.domain.resource_lookup import (
    book_codes_for_lang_from_usfm_only,
    prepare_resource_filepath,
    provision_asset_files,
    resource_lookup_dto,
    resource_types,
)
from doc.utils.docx_util import ensure_reference_styles
from doc.utils.file_utils import docx_filepath, file_needs_update
from doc.utils.text_utils import maybe_correct_book_name
from docx import Document
from docx.oxml import parse_xml
from docx.shared import Length, Mm, Pt, RGBColor
from docx.table import _Cell, _Row
from html4docx import HtmlToDocx  # type: ignore
from passages.domain.model import (
    BibleReferenceWithAvailability,
    Passage,
)
from passages.domain.parser import verse_text_html
from passages.domain.stet_verse_list_parser import BOOK_INDEX, parse_bible_blocks
from passages.reviewers_guide.model import BibleReference
from passages.utils.docx_utils import add_footer, add_header
from pydantic import Json
from passages.domain.parser import clean_chapter_html

if TYPE_CHECKING:
    from typing import TypeAlias

    Cell: TypeAlias = _Cell
    Row: TypeAlias = _Row
else:
    Cell = _Cell
    Row = _Row

logger = settings.logger(__name__)

UNAVAILABLE_COLOR = RGBColor(200, 200, 200)


def format_reference_suffix(reference: BibleReference) -> str:
    if (
        reference.end_chapter
        and reference.end_chapter > 0
        and reference.end_chapter_verse_ref
    ):
        return (
            f"{reference.start_chapter}:"
            f"{reference.start_chapter_verse_ref}-"
            f"{reference.end_chapter}:"
            f"{reference.end_chapter_verse_ref}"
        )
    return f"{reference.start_chapter}:{reference.start_chapter_verse_ref}"


def get_passages(
    bible_references_with_availability: list[BibleReferenceWithAvailability],
    usfm_resource_type: str,
    usfm_books: list[USFMBook],
    resource_type_codes_and_names: Mapping[
        str, str
    ] = settings.RESOURCE_TYPE_CODES_AND_NAMES,
) -> list[Passage]:
    passages: list[Passage] = []
    if not usfm_resource_type:
        resource_type_name = ""
    else:
        resource_type_name = resource_type_codes_and_names[usfm_resource_type]
    usfm_book_index: dict[tuple[str, str], USFMBook] = {
        (b.book_code, b.resource_type_name): b for b in usfm_books
    }
    for bible_reference_with_availability in bible_references_with_availability:
        reference = bible_reference_with_availability.reference
        selected_usfm_book = usfm_book_index.get(
            (reference.book_code, resource_type_name)
        )
        verse_text_html_ = (
            verse_text_html(reference, selected_usfm_book) if selected_usfm_book else ""
        )
        passage = Passage(
            reference=reference,
            localized_reference=f"{reference.book_name} {format_reference_suffix(reference)}",
            passage_text=verse_text_html_,
            is_available=bible_reference_with_availability.is_available,
        )
        passages.append(passage)
    return passages


def get_usfm_books_and_usfm_resource_type(
    bible_references_with_availability: list[BibleReferenceWithAvailability],
    lang_code: str,
    usfm_resource_types: Sequence[str] = settings.USFM_RESOURCE_TYPES,
    preferred_resources: Mapping[str, str] = settings.PREFERRED_RESOURCES,
) -> tuple[list[USFMBook], str]:
    # Invariant: book codes are only those that were available from USFM resources
    book_codes = list(
        dict.fromkeys(
            ref.reference.book_code for ref in bible_references_with_availability
        )
    )
    resource_types_ = resource_types(lang_code, ",".join(book_codes))
    resource_types_codes = list(
        {lang_resource_type_tuple[0] for lang_resource_type_tuple in resource_types_}
    )
    available_usfm_resource_types = [
        resource_type_
        for resource_type_ in resource_types_codes
        if resource_type_ in usfm_resource_types
    ]
    ulb_usfm_resource_types = [
        usfm_resource_type_
        for usfm_resource_type_ in available_usfm_resource_types
        if "ulb" in usfm_resource_type_
    ]
    # Preference logic:
    # 1. Use mapping entry if configured for lang_code and available
    # 2. Otherwise prefer ULB
    # 3. Otherwise fallback to any available USFM resource type
    usfm_resource_type = ""
    preferred_type = preferred_resources.get(lang_code)
    if preferred_type and preferred_type in available_usfm_resource_types:
        usfm_resource_type = preferred_type
    elif ulb_usfm_resource_types:
        usfm_resource_type = ulb_usfm_resource_types[0]
    elif available_usfm_resource_types:
        usfm_resource_type = available_usfm_resource_types[0]
    usfm_books = []
    if usfm_resource_type:
        usfm_book = None
        for book_code in book_codes:
            current_task.update_state(state="Locating assets")
            resource_lookup_dto_ = resource_lookup_dto(
                lang_code, usfm_resource_type, book_code
            )
            if resource_lookup_dto_ and resource_lookup_dto_.url:
                current_task.update_state(state="Provisioning asset files")
                resource_dir = prepare_resource_filepath(resource_lookup_dto_)
                provision_asset_files(resource_lookup_dto_.url, resource_dir)
                current_task.update_state(state="Parsing asset files")
                usfm_book = usfm_book_content(
                    resource_lookup_dto_,
                    resource_dir,
                    False,
                )
                for chapter_num_, chapter_ in usfm_book.chapters.items():
                    cleaned_chapter = clean_chapter_html(chapter_.content)
                    usfm_book.chapters[chapter_num_].verses = split_into_verses(
                        cleaned_chapter
                    )
                usfm_books.append(usfm_book)
    return usfm_books, usfm_resource_type


def generate_docx_document(
    lang0_code: str,
    lang0_name: str,
    lang1_code: Optional[str],
    lang1_name: Optional[str],
    bible_references_with_availability: list[BibleReferenceWithAvailability],
    document_request_key_: str,
    docx_filepath_: str,
) -> str:
    """Generate the content for the Passages document"""
    bible_references_with_availability_lang0: list[BibleReferenceWithAvailability] = [
        b
        for b in bible_references_with_availability
        if b.reference.lang_code == lang0_code
    ]
    usfm_books_lang0, usfm_resource_type_lang0 = get_usfm_books_and_usfm_resource_type(
        bible_references_with_availability_lang0, lang0_code
    )
    bible_references_with_availability_lang1: list[BibleReferenceWithAvailability] = (
        [
            b
            for b in bible_references_with_availability
            if b.reference.lang_code == lang1_code
        ]
        if lang1_code
        else []
    )
    usfm_books_lang1: list[USFMBook] = []
    usfm_resource_type_lang1 = ""
    if lang1_code:
        usfm_books_lang1, usfm_resource_type_lang1 = (
            get_usfm_books_and_usfm_resource_type(
                bible_references_with_availability_lang1, lang1_code
            )
        )
    current_task.update_state(state="Assembling content")
    lang0_passages = get_passages(
        bible_references_with_availability_lang0,
        usfm_resource_type_lang0,
        usfm_books_lang0,
    )
    lang1_passages = []
    if lang1_code:
        lang1_passages = get_passages(
            bible_references_with_availability_lang1,
            usfm_resource_type_lang1,
            usfm_books_lang1,
        )
    current_task.update_state(state="Converting to Docx")
    generate_docx(
        lang0_passages,
        lang1_passages,
        docx_filepath_,
        lang0_code,
        lang0_name,
        lang1_code,
        lang1_name,
    )
    return docx_filepath_


def generate_docx(
    lang0_passages: list[Passage],
    lang1_passages: list[Passage],
    docx_filepath: str,
    lang0_code: str,
    lang0_name: str,
    lang1_code: Optional[str],
    lang1_name: Optional[str],
    show_notes_column: bool = False,
    available_reference_style_name: str = "AvailableReference",
    unavailable_reference_style_name: str = "UnavailableReference",
    unavailable_color: RGBColor = UNAVAILABLE_COLOR,
    document_margin_width: Length = Pt(54),
    a4_width: Length = Mm(210),
    a4_height: Length = Mm(297),
) -> None:
    doc = Document()
    section = doc.sections[0]
    # Explicitly set A4 paper size and uniform margins
    section.page_width = a4_width
    section.page_height = a4_height
    section.left_margin = document_margin_width
    section.right_margin = document_margin_width
    section.top_margin = document_margin_width
    section.bottom_margin = document_margin_width
    # Calculate exact printable area dynamically from margins (EMU level precision)
    printable_width_emu: int = int(a4_width) - (2 * int(document_margin_width))
    ensure_reference_styles(doc, unavailable_color=unavailable_color)
    html_to_docx = HtmlToDocx()
    has_lang1 = lang1_code is not None and lang1_name is not None
    if has_lang1:
        assert len(lang0_passages) == len(lang1_passages), (
            f"Passage count mismatch: "
            f"{len(lang0_passages)} vs {len(lang1_passages)}"
        )
    columns: list[str] = ["lang0"]
    if has_lang1:
        columns.append("lang1")
    if show_notes_column:
        columns.append("notes")
    # Proportionally split printable width so columns perfectly fill the margin bounds
    if columns == ["lang0"]:
        col_widths = [Length(printable_width_emu)]
    elif columns == ["lang0", "lang1"]:
        w = Length(int(printable_width_emu * 0.50))
        col_widths = [w, w]
    elif columns == ["lang0", "notes"]:
        col_widths = [
            Length(int(printable_width_emu * 0.65)),
            Length(int(printable_width_emu * 0.35)),
        ]
    elif columns == ["lang0", "lang1", "notes"]:
        col_widths = [
            Length(int(printable_width_emu * 0.425)),
            Length(int(printable_width_emu * 0.425)),
            Length(int(printable_width_emu * 0.15)),
        ]
    else:
        logger.warning(f"Unexpected column configuration: {columns}")
        col_widths = [Length(int(printable_width_emu / len(columns)))] * len(columns)
    table = doc.add_table(rows=0, cols=len(columns))
    table.autofit = False
    table.allow_autofit = False
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    col_index = {name: i for i, name in enumerate(columns)}
    pairs = (
        zip(lang0_passages, lang1_passages)
        if has_lang1
        else ((p, None) for p in lang0_passages)
    )
    for p0, p1 in pairs:
        row = table.add_row()
        # Apply explicit column widths to cells in each row to prevent Word layout collapse
        for i, w in enumerate(col_widths):
            row.cells[i].width = w
        cell = row.cells[col_index["lang0"]]
        run = cell.add_paragraph().add_run(p0.localized_reference)
        run.style = (
            available_reference_style_name
            if p0.is_available
            else unavailable_reference_style_name
        )
        if p0.is_available:
            html_to_docx.add_html_to_document(p0.passage_text, cell)
        if has_lang1 and p1 is not None:
            cell = row.cells[col_index["lang1"]]
            run = cell.add_paragraph().add_run(p1.localized_reference)
            run.style = (
                available_reference_style_name
                if p1.is_available
                else unavailable_reference_style_name
            )
            if p1.is_available:
                html_to_docx.add_html_to_document(p1.passage_text, cell)
        if show_notes_column:
            cell = row.cells[col_index["notes"]]
            cell.text = ""
            add_vertical_line(cell)
    doc = add_footer(doc)
    doc = add_header(doc, lang0_name, lang1_name, header_text="Passages")
    doc.save(docx_filepath)


# Define the WordprocessingML namespace
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def add_vertical_line(cell: Cell) -> None:
    """Adds a vertical line to the left side of the given table cell."""
    cell_xml = cell._tc  # Access the underlying XML of the cell
    # Ensure the `<w:tcPr>` (table cell properties) element exists
    tc_pr = cell_xml.find(f"{{{WORD_NAMESPACE}}}tcPr")
    if tc_pr is None:
        tc_pr = parse_xml(f'<w:tcPr xmlns:w="{WORD_NAMESPACE}"/>')
        cell_xml.append(tc_pr)
    # Add the left border
    borders_xml = f"""
    <w:tcBorders xmlns:w="{WORD_NAMESPACE}">
        <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
    </w:tcBorders>
    """
    tc_pr.append(parse_xml(borders_xml))


def document_request_key(
    lang0_code: str,
    lang1_code: Optional[str],
    passage_reference_dtos: list[BibleReferenceWithAvailability],
    max_filename_len: int = 240,
    underscore: str = "_",
    hyphen: str = "-",
) -> str:
    """
    Create and return the document_request_key. The
    document_request_key uniquely identifies a document request.

    If the document request key is max_filename_len or more characters
    in length, then switch to using a shorter string that is based on the
    current time. The reason for this is that the document request key is
    used as the file name (with suffix appended) and each OS has a limit
    to how long a file name may be. max_filename_len should make room for
    the file suffix, e.g., ".docx", to be appended.

    It is really useful to have filenames with semantic meaning and so
    those are preferred when possible, i.e., when the file name is not
    too long.
    """

    translation_table = str.maketrans(":;,-", "____")
    passages_key = underscore.join(
        [
            f"{passage_reference.reference.book_code}_{passage_reference.reference.start_chapter}_{passage_reference.reference.start_chapter_verse_ref.translate(translation_table)}"
            for passage_reference in passage_reference_dtos
        ]
    )
    document_request_key_ = (
        f"{lang0_code}_{lang1_code}_{passages_key}_passages"
        if lang1_code
        else f"{lang0_code}_{passages_key}_passages"
    )
    if len(document_request_key_) >= max_filename_len:
        # Likely the generated filename was too long for the OS where this is
        # running. In that case, use the current time as a document_request_key
        # value as doing so results in an acceptably short length.
        timestamp_components = str(time.time()).split(".")
        return f"{timestamp_components[0]}_{timestamp_components[1]}"
    else:
        # Use the semantic filename which declaratively describes the
        # document request components.
        return document_request_key_


@worker.app.task
def generate_passages_docx_document(
    lang0_code: str,
    lang0_name: str,
    lang1_code: Optional[str],
    lang1_name: Optional[str],
    passage_reference_dtos_json: str,
    email_address: str,
    docx_filepath_prefix: str = "passages_",
) -> Json[str]:
    passage_reference_dtos_list = json.loads(passage_reference_dtos_json)
    passage_reference_dtos = [
        BibleReferenceWithAvailability(**d) for d in passage_reference_dtos_list
    ]
    # logger.debug(
    #     "passed args: lang0_code: %s, lang1_code: %s, passage_references: %s, email_adress: %s",
    #     lang0_code,
    #     lang1_code,
    #     passage_reference_dtos,
    #     email_address,
    # )
    document_request_key_ = document_request_key(
        lang0_code, lang1_code, passage_reference_dtos
    )
    docx_filepath_ = f"{docx_filepath(document_request_key_, docx_filepath_prefix)}"
    logger.debug("docx_filepath_: %s", docx_filepath_)
    if file_needs_update(docx_filepath_):
        generate_docx_document(
            lang0_code,
            lang0_name,
            lang1_code,
            lang1_name,
            passage_reference_dtos,
            document_request_key_,
            docx_filepath_,
        )
        if should_send_email(email_address):
            attachments = [
                Attachment(
                    filepath=docx_filepath_,
                    mime_type=(
                        "application",
                        "vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                )
            ]
            current_task.update_state(state="Sending email")
            send_email_with_attachment(
                email_address,
                attachments,
                document_request_key_,
            )
    else:
        logger.debug("Cache hit for %s", docx_filepath_)
    return document_request_key_


@worker.app.task
def stet_exhaustive_verse_list(
    lang_code: str = "en",
    filepath: str = "backend/passages/data/Spiritual_Terms_Evaluation_Exhaustive_Verse_List.txt",
    book_index: dict[str, int] = BOOK_INDEX,
) -> Sequence[BibleReference]:
    """
    >>> from passages.domain.document_generator import stet_exhaustive_verse_list
    >>> result = stet_exhaustive_verse_list()
    >>> result[0]
    Matthew 10:1
    """
    bible_references = []
    with open(filepath, "r") as fi:
        text = fi.read()
        parsed = parse_bible_blocks(text)
        for sublist in parsed.values():
            for value in sublist:
                if len(value.split()) >= 2:
                    bible_references.append(parse_bible_reference(value))
    # Remove duplicates and sort based on the book index, chapter, and verse
    unique_bible_references = sorted(
        set(bible_references),
        key=lambda ref: (
            book_index[ref.book_code],
            ref.start_chapter,
            ref.start_chapter_verse_ref,
        ),
    )
    # Localize the book names
    book_name_map = {
        book_code_and_name[0]: book_code_and_name[1]
        for book_code_and_name in book_codes_for_lang_from_usfm_only(lang_code)
    }
    for bible_reference in unique_bible_references:
        maybe_localized_book_name = book_name_map.get(
            bible_reference.book_code, bible_reference.book_name
        )
        logger.debug("maybe_localized_book_name: %s", maybe_localized_book_name)
        localized_book_name = maybe_correct_book_name(
            lang_code, maybe_localized_book_name
        )
        bible_reference.book_name = localized_book_name
    return unique_bible_references


def get_book_code(book_name: str) -> str:
    return next(code for code, name in BOOK_NAMES.items() if name == book_name)


def parse_bible_reference(book_and_reference_raw: str) -> BibleReference:
    book_name_and_reference = book_and_reference_raw.split()
    book_name = (
        " ".join(book_name_and_reference[:-1])
        if len(book_name_and_reference) > 2
        else book_name_and_reference[0]
    )
    chapter_reference = book_name_and_reference[-1]
    chapter = int(chapter_reference.split(":")[0])
    chapter_verse_ref = chapter_reference.split(":")[1]
    bible_reference = BibleReference(
        lang_code=None,
        book_code=get_book_code(book_name),
        book_name=book_name,
        start_chapter=chapter,
        start_chapter_verse_ref=chapter_verse_ref,
        end_chapter=None,
        end_chapter_verse_ref=None,
    )
    return bible_reference
