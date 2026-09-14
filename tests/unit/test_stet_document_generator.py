import tempfile
from typing import Optional

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from stet.domain.document_generator import generate_docx
from stet.domain.model import VerseEntry, WordEntry


def _generate_test_docx(filepath: str) -> None:
    word_entries = [
        WordEntry(
            words=["word"],
            strongs_numbers="H1234",
            definition="A definition.",
            verses=[
                VerseEntry(
                    source_reference="Gen 1:1",
                    source_text="In the beginning God created the heavens and the earth.",
                    target_reference="Gen 1:1",
                    target_text="Target text for the verse that spans multiple lines.",
                )
            ],
        )
    ]
    generate_docx(word_entries, filepath, lang0_code="en", lang1_code="en")


def _paragraph_line_spacing_twips(paragraph: Paragraph) -> Optional[str]:
    pPr = paragraph._p.find(qn("w:pPr"))
    if pPr is None:
        return None
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        return None
    return spacing.get(qn("w:line"))


def test_table_cell_verse_paragraphs_use_normal_line_spacing() -> None:
    """
    Regression test: table-cell paragraphs holding verse source/target text
    must not carry an explicit doubled line spacing (w:line=480, i.e. 2.0x).
    They should match the document's normal (unset/default) paragraph spacing.
    """
    with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
        _generate_test_docx(tmp.name)
        doc = Document(tmp.name)
        table = doc.tables[0]
        # Row 0 is the header row, row 1 is the reference row, row 2 is the
        # source/target text row that previously had inflated line spacing.
        text_row = table.rows[2]
        source_paragraph = text_row.cells[0].paragraphs[0]
        target_paragraph = text_row.cells[1].paragraphs[0]

        assert source_paragraph.paragraph_format.line_spacing is None
        assert target_paragraph.paragraph_format.line_spacing is None

        assert _paragraph_line_spacing_twips(source_paragraph) != "480"
        assert _paragraph_line_spacing_twips(target_paragraph) != "480"
