from pathlib import Path
from typing import Optional

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import _Cell, _Row
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from passages.domain.document_generator import generate_docx
from passages.domain.model import Passage
from passages.reviewers_guide.model import BibleReference


def _passage(reference: str, html: str, is_available: bool = True) -> Passage:
    return Passage(
        reference=BibleReference(
            lang_code="en",
            book_code="gen",
            book_name="Genesis",
            start_chapter=1,
            start_chapter_verse_ref="1",
            end_chapter=None,
            end_chapter_verse_ref=None,
        ),
        passage_text=html,
        localized_reference=reference,
        is_available=is_available,
    )


def _verse(n: int) -> str:
    return (
        f'<span class="verse"><sup class="versemarker">{n}</sup>Verse {n} text. </span>'
    )


def _poetry(n_lines: int) -> str:
    lines = "".join(
        f'<div class="poetry-1">Line {i} of poem</div>' for i in range(n_lines)
    )
    return f'<span class="verse"><sup class="versemarker">1</sup>{lines}</span>'


def _generate(
    tmp_path: Path,
    lang0: list[Passage],
    lang1: Optional[list[Passage]] = None,
) -> DocxDocument:
    path = str(tmp_path / "out.docx")
    generate_docx(
        lang0,
        lang1 or [],
        path,
        "en",
        "English",
        "fr" if lang1 else None,
        "French" if lang1 else None,
    )
    return Document(path)


def _has(paragraph: Paragraph, tag: str) -> bool:
    return paragraph._p.pPr is not None and paragraph._p.pPr.find(qn(tag)) is not None


def _cant_split(row: _Row) -> bool:
    tr_pr = row._tr.find(qn("w:trPr"))
    return tr_pr is not None and tr_pr.find(qn("w:cantSplit")) is not None


def _passage_paragraphs(cell: _Cell, reference: str) -> list[Paragraph]:
    """Paragraphs from the reference line to the end of the cell"""
    texts = [p.text for p in cell.paragraphs]
    paragraphs: list[Paragraph] = cell.paragraphs
    return paragraphs[texts.index(reference) :]


def _assert_kept(cell: _Cell, reference: str) -> None:
    paragraphs = _passage_paragraphs(cell, reference)
    assert paragraphs[0].text == reference
    for p in paragraphs:
        assert _has(p, "w:keepLines")
    for p in paragraphs[:-1]:
        assert _has(p, "w:keepNext")
    assert not _has(paragraphs[-1], "w:keepNext")


def test_single_language_multi_paragraph_passage(tmp_path: Path) -> None:
    doc = _generate(tmp_path, [_passage("Psalm 1:1-2", _poetry(3))])
    row = doc.tables[0].rows[0]
    assert _cant_split(row)
    cell = row.cells[0]
    _assert_kept(cell, "Psalm 1:1-2")
    # The reference line is glued to the passage text
    assert len(_passage_paragraphs(cell, "Psalm 1:1-2")) >= 4


def test_leading_empty_paragraph_untouched(tmp_path: Path) -> None:
    doc = _generate(tmp_path, [_passage("Gen 1:1", _verse(1))])
    cell = doc.tables[0].rows[0].cells[0]
    assert cell.paragraphs[0].text == ""
    assert not _has(cell.paragraphs[0], "w:keepLines")
    assert not _has(cell.paragraphs[0], "w:keepNext")


def test_short_multi_passage_rows_each_kept_but_not_chained(tmp_path: Path) -> None:
    """Several short passages: each row is kept whole, but nothing glues
    one passage to the next, so several can share a page."""
    passages = [_passage(f"Gen 1:{n}", _verse(n)) for n in range(1, 6)]
    doc = _generate(tmp_path, passages)
    rows = doc.tables[0].rows
    assert len(rows) == 5
    for n, row in enumerate(rows, start=1):
        assert _cant_split(row)
        _assert_kept(row.cells[0], f"Gen 1:{n}")
        # Last paragraph of the passage must not be glued to the next row
        assert not _has(row.cells[0].paragraphs[-1], "w:keepNext")


def test_two_language_rows_track_each_cell_separately(tmp_path: Path) -> None:
    lang0 = [_passage("Psalm 1:1-2", _poetry(2)), _passage("Gen 1:1", _verse(1))]
    lang1 = [_passage("Psaume 1:1-2", _poetry(4)), _passage("Gen 1:1 fr", _verse(1))]
    doc = _generate(tmp_path, lang0, lang1)
    rows = doc.tables[0].rows
    for row, (ref0, ref1) in zip(
        rows, [("Psalm 1:1-2", "Psaume 1:1-2"), ("Gen 1:1", "Gen 1:1 fr")]
    ):
        assert _cant_split(row)
        _assert_kept(row.cells[0], ref0)
        _assert_kept(row.cells[1], ref1)


def test_unavailable_passage_reference_only(tmp_path: Path) -> None:
    doc = _generate(tmp_path, [_passage("Gen 1:1", "", is_available=False)])
    row = doc.tables[0].rows[0]
    assert _cant_split(row)
    paragraphs = _passage_paragraphs(row.cells[0], "Gen 1:1")
    assert len(paragraphs) == 1
    assert _has(paragraphs[0], "w:keepLines")
    assert not _has(paragraphs[0], "w:keepNext")


def test_passage_longer_than_a_page_is_not_cant_split(tmp_path: Path) -> None:
    """A cantSplit row taller than a page is clipped by Word, cutting off
    the end of the passage. A passage taller than a page must therefore
    stay splittable, while a short one next to it is still kept whole."""
    long_passage = _passage("Exodus 14:1-31", _poetry(400))
    short_passage = _passage("Gen 1:1", _verse(1))
    doc = _generate(tmp_path, [long_passage, short_passage])
    rows = doc.tables[0].rows
    assert len(rows) == 2
    assert not _cant_split(rows[0])
    assert _cant_split(rows[1])
    for row, ref in zip(rows, ["Exodus 14:1-31", "Gen 1:1"]):
        _assert_kept(row.cells[0], ref)
    assert len(_passage_paragraphs(rows[0].cells[0], "Exodus 14:1-31")) > 400


def test_long_single_paragraph_passage_is_not_cant_split(tmp_path: Path) -> None:
    """Prose is one paragraph, so tallness comes from line wrapping"""
    html = "".join(_verse(n) for n in range(1, 400))
    doc = _generate(tmp_path, [_passage("Exodus 32:1-35", html)])
    assert not _cant_split(doc.tables[0].rows[0])


def test_tall_passage_in_either_language_leaves_row_splittable(
    tmp_path: Path,
) -> None:
    lang0 = [_passage("Gen 1:1", _verse(1))]
    lang1 = [_passage("Gen 1:1 fr", _poetry(400))]
    doc = _generate(tmp_path, lang0, lang1)
    assert not _cant_split(doc.tables[0].rows[0])


def _prose(n_chars: int) -> str:
    return f'<span class="verse"><sup class="versemarker">1</sup>{"word " * (n_chars // 5)}</span>'


def test_passage_filling_most_of_a_page_stays_cant_split(tmp_path: Path) -> None:
    """A passage that fits on a page by itself, e.g., a page-sized one
    following another passage, must be kept whole so that it moves to the
    next page instead of being split across the page break."""
    previous = _passage("Genesis 2:4-25", _prose(2000))
    # Just under a page by the estimate
    fits = _passage("Genesis 3:1-24", _prose(3800))
    doc = _generate(tmp_path, [previous, fits])
    for row in doc.tables[0].rows:
        assert _cant_split(row)


def test_page_filling_passage_in_two_language_row_stays_cant_split(
    tmp_path: Path,
) -> None:
    """Half-width columns wrap more, so the same page fraction holds less text"""
    lang0 = [_passage("Genesis 2:4-25", _prose(1000)), _passage("Gen 3", _prose(1850))]
    lang1 = [_passage("Ukubala 2:4-25", _prose(1000)), _passage("Uku 3", _prose(1850))]
    doc = _generate(tmp_path, lang0, lang1)
    for row in doc.tables[0].rows:
        assert _cant_split(row)


def _page_break_before(cell: _Cell) -> bool:
    return _has(cell.paragraphs[0], "w:pageBreakBefore")


def test_over_a_page_passage_after_short_one_starts_on_fresh_page(
    tmp_path: Path,
) -> None:
    short = _passage("Ukufuma 12:21-42", _prose(800))
    long = _passage("Ukufuma 14:1-31", _prose(8000))
    doc = _generate(tmp_path, [short, long])
    first, second = doc.tables[0].rows
    assert _cant_split(first)
    assert not _page_break_before(first.cells[0])
    assert not _cant_split(second)
    assert _page_break_before(second.cells[0])
    _assert_kept(second.cells[0], "Ukufuma 14:1-31")


def test_over_a_page_passage_in_either_language_breaks_before_both_cells(
    tmp_path: Path,
) -> None:
    lang0 = [_passage("A", _prose(300)), _passage("B", _prose(300))]
    lang1 = [_passage("A fr", _prose(300)), _passage("B fr", _poetry(400))]
    doc = _generate(tmp_path, lang0, lang1)
    second = doc.tables[0].rows[1]
    assert not _cant_split(second)
    assert _page_break_before(second.cells[0])
    assert _page_break_before(second.cells[1])


def test_over_a_page_first_passage_has_no_page_break_before(tmp_path: Path) -> None:
    doc = _generate(tmp_path, [_passage("Exodus 14:1-31", _prose(8000))])
    row = doc.tables[0].rows[0]
    assert not _cant_split(row)
    assert not _page_break_before(row.cells[0])


def test_short_passages_have_no_page_break_before(tmp_path: Path) -> None:
    doc = _generate(tmp_path, [_passage("A", _verse(1)), _passage("B", _verse(1))])
    for row in doc.tables[0].rows:
        assert not _page_break_before(row.cells[0])


def test_passage_not_fitting_after_near_full_page_passage_starts_fresh_page(
    tmp_path: Path,
) -> None:
    """Near-full page, then a roughly-a-page passage, then a more-than-a-page
    one: neither may start mid-page after the first. The first two fit on a
    page so cantSplit moves them whole; the third is splittable but forces a
    page break before itself."""
    near_full = _passage("Ukufuma 12:21-42", _prose(3800))
    about_a_page = _passage("Ukufuma 14:1-31", _prose(3800))
    more_than_a_page = _passage("Ukufuma 20:1-30", _prose(8000))
    doc = _generate(tmp_path, [near_full, about_a_page, more_than_a_page])
    first, second, third = doc.tables[0].rows
    assert _cant_split(first) and _cant_split(second)
    assert not _page_break_before(first.cells[0])
    assert not _cant_split(third)
    assert _page_break_before(third.cells[0])
    # Nothing glues one passage to the next
    for row, ref in zip(
        (first, second, third),
        ("Ukufuma 12:21-42", "Ukufuma 14:1-31", "Ukufuma 20:1-30"),
    ):
        _assert_kept(row.cells[0], ref)
